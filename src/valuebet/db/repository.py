"""Thin persistence layer between domain objects and the ORM."""

from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
import re
from datetime import UTC, datetime, timedelta, timezone
from typing import Any

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..core.models import MarketSnapshot, ValueSignal
from ..logging import get_logger
from .models import (
    Account,
    AccountNote,
    Bet,
    BetAllocation,
    BookmakerLimitEvent,
    Event,
    OddsSnapshot,
    Signal,
)

log = get_logger("db.repository")

# Statuses for which a signal is considered "still live" in the feed.
OPEN_STATUSES = ("detected", "approved")
DEMO_EVENT_IDS = (0, 1, 2, 3, 4)


# ---------------------------------------------------------------------------
# Helpers for market / fixture / score display
# ---------------------------------------------------------------------------


def classify_bet_type(market_type: str | None, selection: str | None) -> str:
    mt = (market_type or "").strip()
    mt_lower = mt.lower()
    sel = (selection or "").lower()
    if "penalty" in mt_lower or "penalties" in mt_lower or "penalty" in sel:
        return "Penalties"
    if "card" in mt_lower or "booking" in mt_lower or "card" in sel:
        return "Cards"
    if "corner" in mt_lower or "corner" in sel:
        return "Corners"
    if mt_lower.startswith("totals_"):
        line = mt.split("_", 1)[1]
        return f"Over/Under ({line})"
    if "over" in mt_lower or "under" in mt_lower or "total" in mt_lower:
        return "Over/Under"
    if mt_lower.startswith("handicap_"):
        line = mt.split("_", 1)[1]
        return f"Handicap ({line})"
    if "handicap" in mt_lower or "asian" in mt_lower or "spread" in mt_lower:
        return "Handicap"
    if "btts" in mt_lower or "both teams" in mt_lower:
        return "Both Teams to Score"
    if any(k in mt_lower for k in ("h2h", "1x2", "winner", "moneyline", "match")):
        return "Match Winner (1X2)"
    return mt.replace("_", " ").title() if mt else "Match Winner"


def infer_event_name(sig: Signal) -> str | None:
    """Return stored event_name, or reconstruct 'Home vs Away' from all_market_odds_json for legacy rows."""
    if getattr(sig, "event_name", None):
        return sig.event_name
    if sig.all_market_odds_json:
        try:
            odds_map = json.loads(sig.all_market_odds_json)
            if isinstance(odds_map, dict):
                teams = [
                    k for k in odds_map.keys()
                    if k.strip().lower() not in ("draw", "tie", "x")
                    and not k.strip().lower().startswith(("over ", "under "))
                ]
                if len(teams) >= 2:
                    # Clean handicap suffixes if present, e.g. "Arsenal (-1.5)" -> "Arsenal"
                    clean = [re.sub(r"\s*\([+-]?\d+(?:\.\d+)?\)\s*$", "", t).strip() for t in teams[:2]]
                    return f"{clean[0]} vs {clean[1]}"
        except Exception:
            pass
    return None


def extract_final_score(bet: Bet | None) -> str | None:
    """Return bet.final_score or parse legacy score notes if present."""
    if bet is None:
        return None
    if getattr(bet, "final_score", None):
        return bet.final_score
    if bet.note and "scores:[" in bet.note:
        try:
            raw_part = bet.note.split("scores:", 1)[1]
            bracket_end = raw_part.find("]")
            if bracket_end != -1:
                parsed = ast.literal_eval(raw_part[: bracket_end + 1])
                if isinstance(parsed, list) and len(parsed) >= 2:
                    s0 = parsed[0].get("score")
                    s1 = parsed[1].get("score")
                    if s0 is not None and s1 is not None:
                        return f"{s0} - {s1}"
        except Exception:
            pass
    return None


def extract_settled_at_iso(bet: Bet | None) -> str | None:
    if bet is None:
        return None
    if getattr(bet, "settled_at", None):
        return bet.settled_at.isoformat()
    if bet.note and "settled_at:" in bet.note:
        try:
            return bet.note.split("settled_at:")[1].split()[0]
        except Exception:
            pass
    return None


def _populate_signal_market_odds(row: Signal, all_market_odds: dict[str, float] | None) -> None:
    if not all_market_odds:
        return
    draw_item = next(
        (item for item in all_market_odds.items() if item[0].strip().lower() in ("draw", "tie", "x")),
        None,
    )
    non_draw_items = [
        item for item in all_market_odds.items() if item[0].strip().lower() not in ("draw", "tie", "x")
    ]
    is_totals = all(k.strip().lower().startswith(("over ", "under ")) for k, _ in non_draw_items) if non_draw_items else False

    ordered_odds: dict[str, float] = {}
    if not is_totals:
        if len(non_draw_items) >= 1:
            ordered_odds[non_draw_items[0][0]] = non_draw_items[0][1]
            row.home_odds = non_draw_items[0][1]
        if draw_item:
            ordered_odds[draw_item[0]] = draw_item[1]
            row.draw_odds = draw_item[1]
        if len(non_draw_items) >= 2:
            ordered_odds[non_draw_items[1][0]] = non_draw_items[1][1]
            row.away_odds = non_draw_items[1][1]
    else:
        row.home_odds = None
        row.draw_odds = None
        row.away_odds = None

    row.all_market_odds_json = json.dumps(ordered_odds or all_market_odds)


# ---------------------------------------------------------------------------
# Odds / Snapshots
# ---------------------------------------------------------------------------


def save_snapshots(session: Session, snapshots: list[MarketSnapshot]) -> int:
    """Persist every quote in every snapshot to the time-series table."""
    rows = 0
    for snap in snapshots:
        for q in snap.quotes:
            session.add(
                OddsSnapshot(
                    captured_at=q.captured_at,
                    event_id=_safe_event_id(snap.event_id),
                    market_type=snap.market_type,
                    source=q.source,
                    selection=q.selection,
                    decimal_odds=q.decimal_odds,
                    lay_odds=q.lay_odds,
                    back_liquidity=q.back_liquidity,
                    lay_liquidity=q.lay_liquidity,
                    total_matched=snap.total_matched,
                )
            )
            rows += 1
    return rows


# ---------------------------------------------------------------------------
# Signals (with 1.5% Value-Change Re-Read Rule)
# ---------------------------------------------------------------------------


def has_open_signal(session: Session, event_id: int, market_type: str, selection: str) -> bool:
    """True if a still-live signal already exists for this event/market/selection."""
    stmt = (
        select(Signal.id)
        .where(
            Signal.event_id == event_id,
            Signal.market_type == market_type,
            Signal.selection == selection,
            Signal.status.in_(OPEN_STATUSES),
        )
        .limit(1)
    )
    return session.scalar(stmt) is not None


def latest_signal_for_selection(
    session: Session, event_id: int, market_type: str, selection: str
) -> Signal | None:
    """Return the most recent non-cancelled signal for (event_id, market_type, selection)."""
    stmt = (
        select(Signal)
        .where(
            Signal.event_id == event_id,
            Signal.market_type == market_type,
            Signal.selection == selection,
            Signal.status != "cancelled",
        )
        .order_by(Signal.detected_at.desc(), Signal.id.desc())
        .limit(1)
    )
    return session.scalar(stmt)


def save_signal(
    session: Session,
    sig: ValueSignal,
    *,
    min_edge_delta: float | None = None,
) -> Signal | None:
    """Persist a signal, enforcing the minimum value-change (edge delta) re-read threshold.

    A previously detected bet (same event_id + market_type + selection) is ONLY re-read
    if its value (`edge`) changed by at least `re_read_min_edge_delta` (default 1.5% = 0.015).
    - If `abs(new_edge - last_edge) < threshold`: skipped (`None`), preventing duplicate spam
      even after a signal transitions to `paper` or `placed`.
    - If `abs(new_edge - last_edge) >= threshold` and the previous signal is still `detected`:
      updates that open signal in place with the fresh odds, edge, and stake.
    - Otherwise, inserts a new `Signal` row.
    """
    event_id = _safe_event_id(sig.event_id)
    delta_threshold = (
        min_edge_delta
        if min_edge_delta is not None
        else get_settings().re_read_min_edge_delta
    )

    prev = latest_signal_for_selection(session, event_id, sig.market_type, sig.selection)
    if prev is not None:
        edge_diff = abs(sig.edge - prev.edge)
        # Use a tiny float epsilon so exact 0.015 (e.g. 0.085 - 0.070) always passes >= 0.015
        if edge_diff + 1e-9 < delta_threshold:
            log.debug(
                "signal_deduped_below_value_delta",
                event_id=event_id,
                selection=sig.selection,
                market_type=sig.market_type,
                prev_edge=round(prev.edge, 4),
                new_edge=round(sig.edge, 4),
                min_delta=delta_threshold,
            )
            return None

        if prev.status == "detected":
            prev.fair_prob = sig.fair_prob
            prev.confirm_prob = sig.confirm_prob
            prev.target_odds = sig.target_odds
            prev.edge = sig.edge
            prev.recommended_stake = sig.recommended_stake
            prev.detected_at = sig.detected_at
            prev.is_live = sig.is_live
            prev.event_start_time = sig.event_start_time or prev.event_start_time
            if sig.event_name:
                prev.event_name = sig.event_name
            _populate_signal_market_odds(prev, sig.all_market_odds)
            session.flush()
            log.info(
                "signal_reread_updated",
                signal_id=prev.id,
                selection=sig.selection,
                edge=round(sig.edge, 4),
                edge_delta=round(edge_diff, 4),
            )
            return prev

    ev_name = sig.event_name
    if not ev_name and sig.all_market_odds:
        teams = [
            k for k in sig.all_market_odds.keys()
            if k.strip().lower() not in ("draw", "tie", "x")
            and not k.strip().lower().startswith(("over ", "under "))
        ]
        if len(teams) >= 2:
            clean = [re.sub(r"\s*\([+-]?\d+(?:\.\d+)?\)\s*$", "", t).strip() for t in teams[:2]]
            ev_name = f"{clean[0]} vs {clean[1]}"

    row = Signal(
        event_id=event_id,
        event_name=ev_name,
        market_type=sig.market_type,
        selection=sig.selection,
        sport=sig.sport.value,
        fair_prob=sig.fair_prob,
        confirm_prob=sig.confirm_prob,
        target_odds=sig.target_odds,
        edge=sig.edge,
        recommended_stake=sig.recommended_stake,
        status="detected",
        detected_at=sig.detected_at,
        is_live=sig.is_live,
        event_start_time=sig.event_start_time,
        variables_complete=True,
    )
    _populate_signal_market_odds(row, sig.all_market_odds)
    session.add(row)
    session.flush()
    log.info(
        "signal_saved",
        signal_id=row.id,
        event_id=event_id,
        event_name=ev_name,
        selection=sig.selection,
        market_type=sig.market_type,
        sport=sig.sport.value,
        target_bookmaker=sig.target_bookmaker,
        target_odds=sig.target_odds,
        edge=round(sig.edge, 4),
        recommended_stake=sig.recommended_stake,
    )
    return row


def open_signals(session: Session, status: str | None = None) -> list[Signal]:
    stmt = select(Signal).order_by(Signal.detected_at.desc(), Signal.id.desc())
    if status:
        stmt = stmt.where(Signal.status == status)
    else:
        stmt = stmt.where(Signal.status.in_(["detected", "approved"]))
    return list(session.scalars(stmt))


def cancel_signal(session: Session, signal_id: int) -> Signal | None:
    """Cancel a signal (and void its bet if one exists)."""
    sig = session.get(Signal, signal_id)
    if not sig:
        return None
    prev_status = sig.status
    sig.status = "cancelled"
    if sig.bet and sig.bet.outcome == "pending":
        sig.bet.outcome = "void"
        sig.bet.profit = 0.0
    log.info("signal_cancelled", signal_id=signal_id, previous_status=prev_status,
             selection=sig.selection)
    return sig


def update_signal(
    session: Session,
    signal_id: int,
    *,
    recommended_stake: float | None = None,
    max_bet: float | None = None,
    is_live: bool | None = None,
    variables_complete: bool | None = None,
    note: str | None = None,
) -> Signal | None:
    """Patch mutable fields on a signal."""
    sig = session.get(Signal, signal_id)
    if not sig:
        return None
    if recommended_stake is not None:
        sig.recommended_stake = recommended_stake
    if max_bet is not None:
        sig.max_bet = max_bet
    if is_live is not None:
        sig.is_live = is_live
    if variables_complete is not None:
        sig.variables_complete = variables_complete
    if note is not None and sig.bet:
        sig.bet.note = note
    session.flush()
    log.info("signal_updated", signal_id=signal_id, selection=sig.selection)
    return sig


# ---------------------------------------------------------------------------
# Bets
# ---------------------------------------------------------------------------


def record_bet(
    session: Session,
    signal: Signal,
    placed_odds: float,
    stake: float,
    dry_run: bool,
    note: str | None = None,
    requested_stake: float | None = None,
    bookmaker: str = "stoiximan",
    verified: bool = False,
) -> Bet:
    """Persist a placement attempt."""
    actual_edge = (signal.fair_prob * placed_odds) - 1.0 if signal.fair_prob else None
    _requested = requested_stake if requested_stake is not None else stake
    bet = Bet(
        signal_id=signal.id,
        book=bookmaker,
        selection=signal.selection,
        placed_odds=placed_odds,
        requested_stake=_requested,
        stake=stake,
        actual_edge=actual_edge,
        outcome="pending",
        dry_run=dry_run,
        note=note,
        placed_at=datetime.now(UTC),
    )
    if dry_run:
        signal.status = "paper"
    elif verified:
        signal.status = "placed"
    else:
        signal.status = "unconfirmed"
    session.add(bet)
    session.flush()

    _allocate_bet_to_accounts(session, signal, bet)

    log.info(
        "bet_recorded",
        bet_id=bet.id,
        signal_id=signal.id,
        signal_status=signal.status,
        bookmaker=bookmaker,
        selection=signal.selection,
        placed_odds=placed_odds,
        requested_stake=_requested,
        accepted_stake=stake,
        dry_run=dry_run,
        verified_on_platform=verified,
        actual_edge=round(actual_edge, 4) if actual_edge is not None else None,
        note=note,
    )
    if not dry_run and not verified:
        log.error(
            "bet_recorded_unverified",
            bet_id=bet.id,
            bookmaker=bookmaker,
            selection=signal.selection,
            stake=stake,
            msg="bet was NOT found at the bookmaker — verify manually before trusting P&L",
        )
    return bet


def _allocate_bet_to_accounts(session: Session, signal: Signal, _bet: Bet) -> None:
    """Create BetAllocation rows for every active (non-paused) account."""
    accounts = list(session.scalars(select(Account).where(Account.is_paused.is_(False))))
    for acc in accounts:
        allocated = round(signal.recommended_stake * acc.percentage_share, 2)
        session.add(
            BetAllocation(
                signal_id=signal.id,
                account_id=acc.id,
                allocated_stake=allocated,
            )
        )
    if accounts:
        session.flush()


# ---------------------------------------------------------------------------
# P&L
# ---------------------------------------------------------------------------


def pnl_summary(session: Session) -> dict:
    """Realised P&L over settled bets, plus open exposure."""
    all_bets = list(session.scalars(select(Bet)))
    bets = [b for b in all_bets if not b.dry_run]
    paper = [b for b in all_bets if b.dry_run]
    settled = [b for b in bets if b.outcome in {"won", "lost", "void"} and b.profit is not None]
    realised = sum((b.profit or 0.0) for b in settled)
    staked = sum((b.stake or 0.0) for b in settled)
    open_exposure = sum((b.stake or 0.0) for b in bets if b.outcome == "pending")
    unconfirmed_ids = set(
        session.scalars(select(Signal.id).where(Signal.status == "unconfirmed")).all()
    )
    unconfirmed = [b for b in bets if b.signal_id in unconfirmed_ids]
    roi = (realised / staked) if staked else 0.0
    return {
        "bets_total": len(bets),
        "bets_settled": len(settled),
        "realised_pnl": round(realised, 2),
        "total_staked_settled": round(staked, 2),
        "roi": round(roi, 4),
        "open_exposure": round(open_exposure, 2),
        "paper_bets": len(paper),
        "paper_exposure": round(sum(b.stake for b in paper if b.outcome == "pending"), 2),
        "unconfirmed_bets": len(unconfirmed),
    }


def get_event_exposure(session: Session, event_id: int) -> float:
    """Return the total staked amount for all real bets on a given event."""
    stmt = (
        select(Bet.stake)
        .join(Signal)
        .where(Signal.event_id == event_id)
        .where(Bet.dry_run.is_(False))
        .where(Bet.outcome.in_(("pending", "won", "lost")))
    )
    stakes = session.scalars(stmt).all()
    return sum(stakes) if stakes else 0.0


# ---------------------------------------------------------------------------
# Account Management
# ---------------------------------------------------------------------------


def list_accounts(session: Session) -> list[Account]:
    return list(session.scalars(select(Account).order_by(Account.name)))


def get_account(session: Session, account_id: int) -> Account | None:
    return session.get(Account, account_id)


def create_account(
    session: Session,
    *,
    name: str,
    bookmaker: str = "stoiximan",
    percentage_share: float = 1.0,
    initial_deposit: float = 0.0,
) -> Account:
    acc = Account(
        name=name,
        bookmaker=bookmaker,
        percentage_share=percentage_share,
        initial_deposit=initial_deposit,
        is_paused=False,
    )
    session.add(acc)
    session.flush()
    log.info("account_created", account_id=acc.id, name=name, bookmaker=bookmaker,
             percentage_share=percentage_share)
    return acc


def update_account(
    session: Session,
    account_id: int,
    *,
    name: str | None = None,
    bookmaker: str | None = None,
    percentage_share: float | None = None,
    initial_deposit: float | None = None,
    is_paused: bool | None = None,
) -> Account | None:
    acc = session.get(Account, account_id)
    if not acc:
        return None
    if name is not None:
        acc.name = name
    if bookmaker is not None:
        acc.bookmaker = bookmaker
    if percentage_share is not None:
        acc.percentage_share = percentage_share
    if initial_deposit is not None:
        acc.initial_deposit = initial_deposit
    if is_paused is not None:
        acc.is_paused = is_paused
    session.flush()
    log.info("account_updated", account_id=account_id, is_paused=acc.is_paused)
    return acc


def get_account_activity(session: Session, account_id: int) -> dict:
    """Return bets, net profit/loss, total profit, and activity summary for one account."""
    acc = session.get(Account, account_id)
    if not acc:
        return {}

    alloc_stmt = (
        select(BetAllocation, Signal, Bet)
        .join(Signal, BetAllocation.signal_id == Signal.id)
        .outerjoin(Bet, Bet.signal_id == Signal.id)
        .where(BetAllocation.account_id == account_id)
        .order_by(Signal.detected_at.desc())
    )
    rows = session.execute(alloc_stmt).all()

    bets_placed = []
    net_profit = 0.0
    total_staked = 0.0

    for alloc, signal, bet in rows:
        share = acc.percentage_share
        profit_share = (bet.profit or 0.0) * share if bet else 0.0
        staked_share = alloc.allocated_stake

        bets_placed.append({
            "signal_id": signal.id,
            "bet_id": bet.id if bet else None,
            "event_name": infer_event_name(signal),
            "selection": signal.selection,
            "sport": signal.sport,
            "market_type": signal.market_type,
            "is_live": signal.is_live,
            "edge": round(signal.edge, 4),
            "allocated_stake": staked_share,
            "placed_odds": bet.placed_odds if bet else None,
            "outcome": bet.outcome if bet else "pending",
            "final_score": extract_final_score(bet),
            "profit_share": round(profit_share, 2),
            "event_start_time": signal.event_start_time.isoformat() if signal.event_start_time else None,
            "identified_at": signal.detected_at.isoformat() if signal.detected_at else None,
            "placed_at": bet.placed_at.isoformat() if bet and bet.placed_at else None,
        })
        if bet and bet.outcome in {"won", "lost"}:
            net_profit += profit_share
            total_staked += staked_share

    roi = (net_profit / total_staked) if total_staked else 0.0
    return {
        "account_id": account_id,
        "name": acc.name,
        "initial_deposit": acc.initial_deposit,
        "percentage_share": acc.percentage_share,
        "is_paused": acc.is_paused,
        "total_bets": len(bets_placed),
        "net_profit": round(net_profit, 2),
        "roi": round(roi, 4),
        "bets": bets_placed,
    }


# ---------------------------------------------------------------------------
# Account Notes
# ---------------------------------------------------------------------------


def list_account_notes(session: Session, account_id: int) -> list[AccountNote]:
    stmt = (
        select(AccountNote)
        .where(AccountNote.account_id == account_id)
        .order_by(AccountNote.created_at.desc())
    )
    return list(session.scalars(stmt))


def add_account_note(session: Session, account_id: int, content: str) -> AccountNote | None:
    acc = session.get(Account, account_id)
    if not acc:
        return None
    note = AccountNote(account_id=account_id, content=content, created_at=datetime.now(UTC))
    session.add(note)
    session.flush()
    log.info("account_note_added", account_id=account_id, note_id=note.id)
    return note


def delete_account_note(session: Session, note_id: int) -> bool:
    note = session.get(AccountNote, note_id)
    if not note:
        return False
    session.delete(note)
    session.flush()
    return True


# ---------------------------------------------------------------------------
# Reports & Export
# ---------------------------------------------------------------------------

EXPORT_COLUMN_LABELS: dict[str, str] = {
    "signal_id": "Signal ID",
    "bet_id": "Bet ID",
    "event_name": "Match / Fixture",
    "selection": "Actual Bet Selection",
    "sport": "Sport",
    "bet_type": "Bet Type",
    "market_type": "Market Type",
    "is_live": "Live / Pre-Match",
    "final_score": "Match Score",
    "outcome": "Outcome (Auto)",
    "event_start_time": "Match Played At",
    "target_bookmaker": "Platform (Stoiximan)",
    "target_odds": "Target Odds (Stoiximan)",
    "betfair_odds": "Betfair Exchange Odds",
    "pinnacle_odds": "Pinnacle Odds",
    "edge": "Edge %",
    "fair_prob": "Fair Probability",
    "confirm_prob": "Confirmation Probability",
    "placed_odds": "Odds at Placement",
    "recommended_stake": "Recommended Stake",
    "requested_stake": "Stake Requested",
    "stake": "Stake Accepted",
    "account_name": "Client / Player Name",
    "allocated_stake": "Player Stake Share",
    "player_profit": "Player Net Profit / Loss",
    "potential_profit": "Potential Win Profit",
    "potential_loss": "Potential Loss",
    "profit": "Net Profit / Loss",
    "actual_edge": "Actual Edge %",
    "clv": "Closing Line Value",
    "dry_run": "Paper Bet (Dry Run)",
    "detected_at": "Bet Identified At",
    "placed_at": "Bet Placed At",
    "settled_at": "Result Settled At",
    "status": "Signal Status",
    "max_bet": "Max Bet Available",
    "variables_complete": "All Variables Available",
}


def generate_report(
    session: Session,
    *,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    sport: str | None = None,
    status: str | None = None,
    bookmaker: str | None = None,
    include_dry_run: bool = True,
    account_id: int | None = None,
) -> list[dict[str, Any]]:
    """Query bets+signals with optional filters and return list of dicts."""
    stmt = (
        select(Signal, Bet)
        .outerjoin(Bet, Bet.signal_id == Signal.id)
        .order_by(Signal.detected_at.desc())
    )
    if date_from:
        stmt = stmt.where(Signal.detected_at >= date_from)
    if date_to:
        stmt = stmt.where(Signal.detected_at <= date_to)
    if sport:
        stmt = stmt.where(Signal.sport == sport)
    if status:
        stmt = stmt.where(Signal.status == status)
    if bookmaker and bookmaker != "all":
        stmt = stmt.where(Bet.book == bookmaker)
    if not include_dry_run:
        stmt = stmt.where((Bet.dry_run.is_(False)) | (Bet.id.is_(None)))

    acc = session.get(Account, account_id) if account_id else None

    rows = session.execute(stmt).all()
    results = []
    for sig, bet in rows:
        live_label = "Live" if sig.is_live else "Pre-Match"
        effective_odds = (bet.placed_odds if bet and bet.placed_odds else sig.target_odds)
        effective_stake = (bet.stake if bet and bet.stake else sig.recommended_stake)
        pot_profit_val = round(effective_stake * (effective_odds - 1), 2)
        pot_profit_pct = round((effective_odds - 1) * 100, 1)

        bf_odds = round(1.0 / sig.fair_prob, 2) if sig.fair_prob and sig.fair_prob > 0 else None
        pin_odds = round(1.0 / sig.confirm_prob, 2) if sig.confirm_prob and sig.confirm_prob > 0 else None
        btype = classify_bet_type(sig.market_type, sig.selection)
        ev_name = infer_event_name(sig) or "—"
        score_str = extract_final_score(bet) or "—"
        settled_iso = extract_settled_at_iso(bet)
        settled_fmt = None
        if settled_iso:
            try:
                settled_fmt = datetime.fromisoformat(settled_iso).strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                settled_fmt = settled_iso

        row_dict: dict[str, Any] = {
            "signal_id": sig.id,
            "bet_id": bet.id if bet else None,
            "event_name": ev_name,
            "selection": sig.selection,
            "sport": sig.sport,
            "bet_type": btype,
            "market_type": sig.market_type,
            "is_live": live_label,
            "final_score": score_str,
            "outcome": bet.outcome if bet else "no bet",
            "event_start_time": sig.event_start_time.strftime("%Y-%m-%d %H:%M:%S") if sig.event_start_time else None,
            "target_bookmaker": bet.book if bet and bet.book else "stoiximan",
            "target_odds": sig.target_odds,
            "betfair_odds": bf_odds,
            "pinnacle_odds": pin_odds,
            "edge": f"{sig.edge * 100:+.2f}%",
            "fair_prob": f"{sig.fair_prob * 100:.2f}%",
            "confirm_prob": f"{sig.confirm_prob * 100:.2f}%" if sig.confirm_prob else "—",
            "placed_odds": bet.placed_odds if bet else None,
            "recommended_stake": sig.recommended_stake,
            "requested_stake": bet.requested_stake if bet else None,
            "stake": bet.stake if bet else None,
        }

        if acc:
            client_stake = round(effective_stake * acc.percentage_share, 2)
            client_profit = round(bet.profit * acc.percentage_share, 2) if bet and bet.profit is not None else None
            row_dict["account_name"] = acc.name
            row_dict["allocated_stake"] = client_stake
            row_dict["player_profit"] = client_profit

        row_dict.update({
            "potential_profit": f"+{pot_profit_pct:.0f}% (+€{pot_profit_val:.2f})",
            "potential_loss": f"-100% (-€{effective_stake:.2f})",
            "profit": bet.profit if bet else None,
            "actual_edge": f"{bet.actual_edge * 100:+.2f}%" if bet and bet.actual_edge else None,
            "clv": f"{bet.clv * 100:+.2f}%" if bet and bet.clv else None,
            "dry_run": bet.dry_run if bet else None,
            "detected_at": sig.detected_at.strftime("%Y-%m-%d %H:%M:%S") if sig.detected_at else None,
            "placed_at": bet.placed_at.strftime("%Y-%m-%d %H:%M:%S") if bet and bet.placed_at else None,
            "settled_at": settled_fmt,
            "status": sig.status,
            "max_bet": sig.max_bet,
            "variables_complete": "Yes" if sig.variables_complete else "No",
        })
        results.append(row_dict)
    return results


def export_csv(session: Session, **kwargs: Any) -> str:
    """Return a CSV string with human-friendly column headers."""
    data = generate_report(session, **kwargs)
    if not data:
        return ""
    output = io.StringIO()
    keys = list(data[0].keys())
    friendly_headers = [EXPORT_COLUMN_LABELS.get(k, k) for k in keys]
    writer = csv.DictWriter(output, fieldnames=keys)
    output.write(",".join(f'"{h}"' for h in friendly_headers) + "\n")
    for row in data:
        writer.writerow(row)
    return output.getvalue()


def export_xlsx(session: Session, **kwargs: Any) -> bytes:
    """Return an Excel (.xlsx) byte string with human-friendly column headers."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    data = generate_report(session, **kwargs)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Betting Report"

    if not data:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    keys = list(data[0].keys())
    friendly_headers = [EXPORT_COLUMN_LABELS.get(k, k) for k in keys]

    header_fill = PatternFill(start_color="1a1f2e", end_color="1a1f2e", fill_type="solid")
    header_font = Font(color="a5b4fc", bold=True, name="Calibri")
    for col_idx, header in enumerate(friendly_headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(data, 2):
        for col_idx, key in enumerate(keys, 1):
            val = row.get(key)
            ws.cell(row=row_idx, column=col_idx, value=val)

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# CLV & Auto-Settlement
# ---------------------------------------------------------------------------


def update_clv_for_pending_bets(session: Session) -> int:
    """Find pending bets for events starting within 10 minutes, calculate CLV, and save."""
    from ..core.odds_math import implied_prob, midpoint_prob

    now = datetime.now(UTC)
    target_time = now + timedelta(minutes=10)

    stmt = (
        select(Bet, Signal, Event)
        .join(Signal, Bet.signal_id == Signal.id)
        .join(Event, Signal.event_id == Event.id)
        .where(Bet.outcome == "pending")
        .where(Bet.clv == None)
        .where(Event.start_time <= target_time)
    )

    rows = session.execute(stmt).all()
    updated = 0

    for bet, signal, event in rows:
        snap_stmt = (
            select(OddsSnapshot)
            .where(OddsSnapshot.event_id == event.id)
            .where(OddsSnapshot.source == "betfair")
            .where(OddsSnapshot.selection == signal.selection)
            .order_by(OddsSnapshot.captured_at.desc())
            .limit(1)
        )
        snap = session.scalar(snap_stmt)
        if snap:
            try:
                if snap.lay_odds is not None:
                    closing_prob = midpoint_prob(snap.decimal_odds, snap.lay_odds)
                else:
                    closing_prob = implied_prob(snap.decimal_odds)

                closing_fair_odds = 1.0 / closing_prob if closing_prob > 0 else 0
                if closing_fair_odds > 0:
                    bet.clv = (bet.placed_odds / closing_fair_odds) - 1.0
                    updated += 1
            except ValueError:
                pass

    if updated > 0:
        session.flush()
    log.info("clv_updated", candidates=len(rows), updated=updated)
    return updated


def settle_pending_bets(session: Session, *, force_check_all: bool = False) -> int:
    """Auto-settle pending bets against official match results from The Odds API scores."""
    from ..core.results import BetOutcome, TheOddsAPIResultResolver
    from ..engine.matching import normalise
    from ..sources.the_odds_api import _parse_commence_time

    stmt = (
        select(Bet, Signal)
        .join(Signal, Bet.signal_id == Signal.id)
        .where(Bet.outcome == "pending")
    )
    rows = session.execute(stmt).all()
    if not rows:
        return 0

    now = datetime.now(timezone.utc)
    resolver = TheOddsAPIResultResolver()

    # Group ONLY real, plausibly-completed bets by sport so fake mock rows ("Team A")
    # or matches kicking off tomorrow never burn API quota every 15 minutes.
    bets_by_sport: dict[str, list[tuple[Bet, Signal]]] = {}
    for bet, sig in rows:
        if sig.event_id in DEMO_EVENT_IDS or sig.selection in {"Team A", "Team B", "Team C", "Team D"}:
            continue
        if not force_check_all and sig.event_start_time is not None:
            st = sig.event_start_time
            if st.tzinfo is None:
                st = st.replace(tzinfo=timezone.utc)
            # A pre-match game cannot be completed until at least ~75m after kickoff
            if now < st + timedelta(minutes=75):
                continue
        bets_by_sport.setdefault(sig.sport, []).append((bet, sig))

    if not bets_by_sport:
        return 0

    settled_count = 0
    now_iso = now.isoformat()

    for sport_str, bet_pairs in bets_by_sport.items():
        events = resolver.fetch_scores_for_sport(sport_str)
        if not events:
            continue

        for bet, sig in bet_pairs:
            matched_event = None
            sig_ev_name = infer_event_name(sig)
            sig_ev_norm = normalise(sig_ev_name) if sig_ev_name else None

            for ev in events:
                if not ev.get("completed"):
                    continue
                ev_id = _safe_event_id(ev.get("id", ""))
                home_team = ev.get("home_team", "")
                away_team = ev.get("away_team", "")

                if ev_id == sig.event_id:
                    matched_event = ev
                    break

                # Fallback match ONLY if BOTH teams match AND start time is within 24h
                if sig_ev_norm and home_team and away_team:
                    cand_norm = normalise(f"{home_team} vs {away_team}")
                    if cand_norm == sig_ev_norm:
                        ev_start = _parse_commence_time(ev.get("commence_time"))
                        if sig.event_start_time and ev_start:
                            st = sig.event_start_time
                            if st.tzinfo is None:
                                st = st.replace(tzinfo=timezone.utc)
                            if abs((st - ev_start).total_seconds()) > 86400:
                                continue
                        matched_event = ev
                        break

            if not matched_event:
                continue

            home_team = matched_event.get("home_team", "")
            away_team = matched_event.get("away_team", "")
            scores_list = matched_event.get("scores", [])

            outcome = resolver.determine_match_outcome(
                selection=bet.selection,
                home_team=home_team,
                away_team=away_team,
                scores=scores_list,
                market_type=sig.market_type,
            )

            if outcome in (BetOutcome.WON, BetOutcome.LOST, BetOutcome.VOID):
                bet.outcome = outcome.value
                if outcome == BetOutcome.WON:
                    bet.profit = round(bet.stake * (bet.placed_odds - 1), 2)
                elif outcome == BetOutcome.LOST:
                    bet.profit = -bet.stake
                else:
                    bet.profit = 0.0

                score_text = resolver.format_score(home_team, away_team, scores_list)
                bet.final_score = score_text
                bet.settled_at = now
                if not sig.event_name and home_team and away_team:
                    sig.event_name = f"{home_team} vs {away_team}"

                sig.status = "settled"
                clean_note = bet.note or ""
                bet.note = f"settled_at:{now_iso} auto_settled_by:the_odds_api score:{score_text or 'N/A'} {clean_note}".strip()[:510]
                settled_count += 1
                log.info(
                    "bet_auto_settled",
                    bet_id=bet.id,
                    event_name=sig.event_name,
                    selection=bet.selection,
                    market_type=sig.market_type,
                    score=score_text,
                    outcome=bet.outcome,
                    profit=bet.profit,
                )

    if settled_count > 0:
        session.flush()
    return settled_count


def purge_demo_data(session: Session) -> dict[str, int]:
    """Delete fabricated mock/demo signals, bets, allocations, and snapshots (evt-0..evt-4)."""
    demo_sig_ids = list(
        session.scalars(
            select(Signal.id).where(
                (Signal.event_id.in_(DEMO_EVENT_IDS))
                | (Signal.selection.in_(["Team A", "Team B", "Team C", "Team D"]))
            )
        )
    )
    bets_deleted = 0
    allocs_deleted = 0
    signals_deleted = 0
    if demo_sig_ids:
        allocs_deleted = session.execute(
            delete(BetAllocation).where(BetAllocation.signal_id.in_(demo_sig_ids))
        ).rowcount or 0
        bets_deleted = session.execute(
            delete(Bet).where(Bet.signal_id.in_(demo_sig_ids))
        ).rowcount or 0
        signals_deleted = session.execute(
            delete(Signal).where(Signal.id.in_(demo_sig_ids))
        ).rowcount or 0
    snaps_deleted = session.execute(
        delete(OddsSnapshot).where(OddsSnapshot.event_id.in_(DEMO_EVENT_IDS))
    ).rowcount or 0
    session.flush()
    log.info(
        "demo_data_purged",
        signals_deleted=signals_deleted,
        bets_deleted=bets_deleted,
        allocs_deleted=allocs_deleted,
        snaps_deleted=snaps_deleted,
    )
    return {
        "signals_deleted": signals_deleted,
        "bets_deleted": bets_deleted,
        "allocations_deleted": allocs_deleted,
        "snapshots_deleted": snaps_deleted,
    }


# ---------------------------------------------------------------------------
# Feature 4: Bookmaker Limit Tracking — DB persistence
# ---------------------------------------------------------------------------


def save_limit_event(
    session: Session,
    bookmaker: str,
    requested_stake: float,
    accepted_stake: float,
    was_rejected: bool,
    note: str | None = None,
) -> BookmakerLimitEvent:
    """Persist one stake-acceptance record to the DB."""
    ratio = (accepted_stake / requested_stake) if requested_stake > 0 else 0.0
    event = BookmakerLimitEvent(
        bookmaker=bookmaker,
        requested_stake=requested_stake,
        accepted_stake=accepted_stake,
        acceptance_ratio=round(ratio, 4),
        was_rejected=was_rejected,
        note=note,
        placed_at=datetime.now(UTC),
    )
    session.add(event)
    session.flush()
    log.debug("limit_event_saved", bookmaker=bookmaker, requested=requested_stake,
              accepted=accepted_stake, ratio=event.acceptance_ratio,
              rejected=was_rejected)
    return event


def bookmaker_limit_summary(session: Session, bookmaker: str, last_n: int = 50) -> dict:
    """Query DB for the last N limit events and compute aggregate acceptance stats."""
    stmt = (
        select(BookmakerLimitEvent)
        .where(BookmakerLimitEvent.bookmaker == bookmaker)
        .order_by(BookmakerLimitEvent.placed_at.desc())
        .limit(last_n)
    )
    events = list(session.scalars(stmt))
    if not events:
        return {"bookmaker": bookmaker, "total_events": 0,
                "acceptance_rate": None, "is_likely_limited": False}

    total_requested = sum(e.requested_stake for e in events)
    total_accepted = sum(e.accepted_stake for e in events)
    rate = (total_accepted / total_requested) if total_requested > 0 else 0.0
    rejections = sum(1 for e in events if e.was_rejected)
    return {
        "bookmaker": bookmaker,
        "total_events": len(events),
        "acceptance_rate": round(rate, 4),
        "is_likely_limited": rate < 0.70 and len(events) >= 5,
        "rejection_count": rejections,
        "last_event": events[0].placed_at.isoformat() if events else None,
    }


def all_bookmaker_limit_summaries(session: Session) -> list[dict]:
    """Return limit summaries for all bookmakers that have events."""
    stmt = select(BookmakerLimitEvent.bookmaker).distinct()
    bookmakers = list(session.scalars(stmt))
    return [bookmaker_limit_summary(session, bk) for bk in bookmakers]


# ---------------------------------------------------------------------------
# Deterministic ID Helper
# ---------------------------------------------------------------------------


def _safe_event_id(raw: str | int) -> int:
    """Source event ids are strings; map them to a stable int for storage."""
    if isinstance(raw, int):
        return raw
    digits = "".join(ch for ch in str(raw) if ch.isdigit())
    if digits:
        return int(digits[:18])
    h = hashlib.sha256(str(raw).encode()).hexdigest()
    return int(h[:15], 16)
